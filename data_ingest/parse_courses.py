"""解析 CISP 课件，支持 legacy/fixed/structured 三种切片策略。"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config
from data_ingest.chunking import (
    build_fixed_chunks,
    build_structured_chunks,
    clean_text,
    deduplicate_chunks,
    split_long,
)
from data_ingest.tables import (
    body_text_without_tables,
    extract_page_tables,
    find_caption,
)
from domains import FILENAME_DOMAINS


def domain_from_filename(name: str) -> str:
    for key, domain in FILENAME_DOMAINS.items():
        if key in name:
            return domain
    return "未分类"


def extract_pdf(path: str, kind: str) -> tuple[list[dict], int]:
    """逐页提取文本；检测到合格表格的页面额外输出 body_text（剔除表格区域）
    与 tables（markdown + 表题），供 structured 切片生成独立表格 Child。"""
    import pymupdf

    domain = domain_from_filename(os.path.basename(path))
    source = os.path.splitext(os.path.basename(path))[0]
    pages = []
    table_count = 0
    doc = pymupdf.open(path)
    for page_no, page in enumerate(doc, start=1):
        entry = {
            "domain": domain,
            "source": source,
            "page": page_no,
            "kind": kind,
            "text": page.get_text(),
        }
        if config.ENABLE_TABLE_CHUNKING:
            tables = extract_page_tables(page)
            if tables:
                text_blocks = [b for b in page.get_text("blocks") if b[6] == 0]
                bboxes = [pymupdf.Rect(t["bbox"]) for t in tables]
                entry["body_text"] = body_text_without_tables(text_blocks, bboxes)
                for table in tables:
                    table["caption"] = find_caption(text_blocks, pymupdf.Rect(table["bbox"]))
                entry["tables"] = tables
                table_count += len(tables)
        pages.append(entry)
    doc.close()
    return pages, table_count


def extract_pptx(path: str) -> list[dict]:
    from pptx import Presentation

    domain = domain_from_filename(os.path.basename(path))
    source = os.path.splitext(os.path.basename(path))[0]
    pages = []
    prs = Presentation(path)
    for slide_no, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                line = "".join(run.text for run in paragraph.runs).strip()
                if line:
                    texts.append(line)
        pages.append({
            "domain": domain,
            "source": source,
            "page": slide_no,
            "kind": "pptx",
            "text": "\n".join(texts),
        })
    return pages


def build_legacy_chunks(pages: list[dict], max_chars: int) -> list[dict]:
    chunks = []
    for page in pages:
        text = clean_text(page["text"])
        if len(text) < 10:
            continue
        for part in split_long(text, max_chars):
            chunks.append({
                "domain": page["domain"],
                "source": page["source"],
                "page": page["page"],
                "kind": page["kind"],
                "text": part,
            })
    return chunks


def _documents() -> list[tuple[list[dict], int]]:
    documents: list[tuple[list[dict], int]] = []
    for root, dirs, files in os.walk(config.COURSE_DIR):
        dirs.sort()
        for filename in sorted(files):
            if filename.lower().endswith(".pdf") and "知识点总结" not in filename:
                pages, _ = extract_pdf(os.path.join(root, filename), "courseware")
                documents.append((pages, config.CHUNK_MAX_CHARS))

    for filename in sorted(os.listdir(config.EXAM_DIR)):
        if filename.lower().endswith(".pptx"):
            documents.append((
                extract_pptx(os.path.join(config.EXAM_DIR, filename)),
                config.CHUNK_MAX_CHARS,
            ))

    for root, dirs, files in os.walk(config.COURSE_DIR):
        dirs.sort()
        for filename in sorted(files):
            if filename.lower().endswith(".pdf") and "知识点总结" in filename:
                pages, _ = extract_pdf(os.path.join(root, filename), "summary")
                documents.append((pages, config.SUMMARY_CHUNK_MAX_CHARS))
    return documents


def parse_all(strategy: str) -> tuple[list[dict], int]:
    chunks: list[dict] = []
    for pages, legacy_max_chars in _documents():
        if strategy == "legacy":
            chunks.extend(build_legacy_chunks(pages, legacy_max_chars))
        elif strategy == "fixed":
            chunks.extend(build_fixed_chunks(
                pages,
                target_chars=config.FIXED_CHUNK_TARGET_CHARS,
                max_chars=config.FIXED_CHUNK_MAX_CHARS,
                overlap_chars=config.FIXED_CHUNK_OVERLAP_CHARS,
            ))
        else:
            chunks.extend(build_structured_chunks(
                pages,
                target_chars=config.CHUNK_TARGET_CHARS,
                max_chars=config.STRUCTURED_CHUNK_MAX_CHARS,
                overlap_units=config.CHUNK_OVERLAP_UNITS,
            ))

    if strategy in ("legacy", "fixed"):
        for index, chunk in enumerate(chunks):
            chunk["id"] = f"chunk_{index:05d}"
            chunk["parent_id"] = chunk["id"]
        return chunks, 0
    return deduplicate_chunks(chunks)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--strategy", choices=("legacy", "fixed", "structured"),
        default=config.CHUNK_STRATEGY,
    )
    parser.add_argument("--output", default=config.KB_CHUNKS_PATH)
    args = parser.parse_args()

    chunks, removed_duplicates = parse_all(args.strategy)
    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    with open(args.output, "w", encoding="utf-8") as fp:
        json.dump(chunks, fp, ensure_ascii=False, indent=1)

    if os.path.abspath(args.output) == os.path.abspath(config.KB_CHUNKS_PATH):
        from data_ingest.manifest import KB_SCHEMA_VERSION, update_manifest
        update_manifest(
            chunk_count=len(chunks),
            kb_schema_version=KB_SCHEMA_VERSION,
            chunking_strategy=args.strategy + ("+tables" if config.ENABLE_TABLE_CHUNKING else ""),
        )

    by_domain: dict[str, int] = {}
    by_kind: dict[str, int] = {}
    for chunk in chunks:
        by_domain[chunk["domain"]] = by_domain.get(chunk["domain"], 0) + 1
        by_kind[chunk["kind"]] = by_kind.get(chunk["kind"], 0) + 1
    table_chunks = sum(1 for chunk in chunks if chunk.get("element") == "table")
    print(f"切片策略: {args.strategy}（表格感知: {'开' if config.ENABLE_TABLE_CHUNKING else '关'}）")
    print(f"共 {len(chunks)} 个知识块（其中表格块 {table_chunks} 个）→ {args.output}")
    print(f"去除重复块: {removed_duplicates}")
    print("按知识域:", json.dumps(by_domain, ensure_ascii=False))
    print("按类型:", json.dumps(by_kind, ensure_ascii=False))
    if "未分类" in by_domain:
        print("[警告] 存在未分类文件，请检查 FILENAME_DOMAINS 映射")


if __name__ == "__main__":
    main()
